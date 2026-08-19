"""
BigStep-branded python-docx toolkit — reusable composition helpers.

The single source of truth for rendering a BigStep-branded .docx from a JSON
content spec. Palette / fonts / logo come from bigstep_brand.py; this module adds
the python-docx primitives (font plumbing, branded tables, callouts, bullets,
section headings, cover header/footer) so a build.py only has to walk a spec.

Deterministic + reproducible: same spec in, same document out (standards P-2.2 /
P-4.5 — codify the mechanical transform instead of hand-building docx per run).
"""
import os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bigstep_brand as bs
from bigstep_brand import C, BRAND
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Poppins"


def rgb(h):
    return RGBColor(*bs.hx(h))


# ------------------------------------------------------------------ font plumbing
def _rfonts(rpr, name):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)


def set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    _rfonts(run._element.get_or_add_rPr(), FONT)


def set_doc_default_font(doc, name=FONT):
    el = doc.styles.element
    dd = el.find(qn('w:docDefaults'))
    if dd is None:
        dd = OxmlElement('w:docDefaults'); el.insert(0, dd)
    rprd = dd.find(qn('w:rPrDefault'))
    if rprd is None:
        rprd = OxmlElement('w:rPrDefault'); dd.append(rprd)
    rpr = rprd.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr'); rprd.append(rpr)
    _rfonts(rpr, name)
    n = doc.styles['Normal']; n.font.name = name
    _rfonts(n.element.get_or_add_rPr(), name)


# ------------------------------------------------------------------ cell helpers
def shade(cell, h):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), h.lstrip('#')); tcPr.append(sh)


def _set_border(edge_el, val, sz, color):
    edge_el.set(qn('w:val'), val)
    if val != 'nil':
        edge_el.set(qn('w:sz'), str(sz)); edge_el.set(qn('w:space'), '0')
        edge_el.set(qn('w:color'), color.lstrip('#'))


def cell_borders(cell, top=None, left=None, bottom=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None:
        tcPr.remove(existing)
    b = OxmlElement('w:tcBorders')
    for name, spec in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{name}')
        if spec is None:
            _set_border(e, 'nil', 0, C.WHITE)
        else:
            _set_border(e, 'single', spec[0], spec[1])
        b.append(e)
    tcPr.append(b)


def left_accent(cell, sz, color):
    cell_borders(cell, left=(sz, color))


def cell_pad(cell, dxa=140, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    vals = {'top': top if top is not None else dxa,
            'bottom': bottom if bottom is not None else dxa, 'start': dxa, 'end': dxa}
    for m, v in vals.items():
        e = OxmlElement(f'w:{m}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcPr.append(mar)


def no_borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)


def set_col_widths(table, widths):
    table.autofit = False; table.allow_autofit = False
    for row in table.rows:
        for j, w in enumerate(widths):
            if j < len(row.cells):
                row.cells[j].width = w


# ------------------------------------------------------------------ text helpers
def para(doc, runs, after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    for (t, s, c, b) in runs:
        set_font(p.add_run(t), s, c, b)
    return p


def cell_text(cell, runs, after=0, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.05, first=True):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    for (t, s, c, b) in runs:
        set_font(p.add_run(t), s, c, b)
    return p


def two_tone_title(doc, lead, tail, size=24, before=4, after=2):
    return para(doc, [(lead, size, C.NAVY, True), (tail, size, C.PRIMARY, True)], after=after, before=before)


def accent_rule(doc, width_in=1.4, color=None):
    color = color or C.CYAN
    t = doc.add_table(rows=1, cols=1); set_col_widths(t, [Inches(width_in)])
    t.rows[0].height = Pt(2.2); no_borders(t); shade(t.rows[0].cells[0], color)
    r = t.rows[0].cells[0].paragraphs[0]; r.paragraph_format.space_after = Pt(0)
    set_font(r.add_run(" "), 2, color)
    s = doc.add_paragraph(); s.paragraph_format.space_after = Pt(1)
    s.paragraph_format.space_before = Pt(0); s.paragraph_format.line_spacing = 1.0
    rr = s.add_run(" "); rr.font.size = Pt(2); rr.font.name = FONT
    _rfonts(rr._element.get_or_add_rPr(), FONT)


def spacer(doc, pts=6):
    doc.add_paragraph().paragraph_format.space_after = Pt(pts)


def section_head(doc, number, title, before=8):
    """Numbered H1 section heading: 'N. Title' two-tone + cyan rule."""
    lead = f"{number}. " if number else ""
    two_tone_title(doc, lead, title, size=14, before=before, after=1)
    accent_rule(doc, 1.1)


def subhead(doc, text, before=4):
    para(doc, [(text, 11, C.NAVY, True)], before=before, after=1, line=1.1)


# ------------------------------------------------------------------ content blocks
def bullets(doc, items, level=0):
    for it in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.28 + 0.22 * level); pf.space_after = Pt(2); pf.line_spacing = 1.14
        set_font(p.add_run(("•  " if level == 0 else "–  ")), 10, C.PRIMARY, True)
        set_font(p.add_run(str(it)), 9.5, C.SLATE, False)


def callout(doc, title, body_lines, accent=None, fill=None, width_in=6.9):
    fill = fill or C.TINT; accent = accent or C.PRIMARY
    t = doc.add_table(rows=1, cols=1); set_col_widths(t, [Inches(width_in)]); no_borders(t)
    cell = t.rows[0].cells[0]; shade(cell, fill); left_accent(cell, 34, accent); cell_pad(cell, 150, top=16, bottom=16)
    if title:
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.05
        set_font(p.add_run(title), 10, C.NAVY, True)
        for ln in body_lines:
            pp = cell.add_paragraph(); pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.1
            set_font(pp.add_run(str(ln)), 9, C.SLATE, False)
    else:
        for i, ln in enumerate(body_lines):
            pp = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.1
            set_font(pp.add_run(str(ln)), 9, C.SLATE, False)
    spacer(doc, 3)


def branded_table(doc, headers, rows, widths_in=None, total_in=6.9):
    """Header row (blue) + zebra data rows. `rows` is a list of lists (cells)."""
    ncol = len(headers)
    t = doc.add_table(rows=1 + len(rows), cols=ncol); no_borders(t)
    # header
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]; shade(cell, C.PRIMARY); cell_pad(cell, 110, top=7, bottom=7)
        cell_text(cell, [(str(h), 9.0, C.WHITE, True)], line=1.02)
    # body (zebra)
    for i, row in enumerate(rows):
        fill = C.WHITE if (i % 2 == 0) else C.TINT2
        for j in range(ncol):
            cell = t.rows[i + 1].cells[j]; shade(cell, fill); cell_pad(cell, 110, top=6, bottom=6)
            val = row[j] if j < len(row) else ""
            cell_text(cell, [(str(val), 8.6, C.SLATE, False)], line=1.05)
    if widths_in:
        set_col_widths(t, [Inches(w) for w in widths_in])
    else:
        set_col_widths(t, [Inches(total_in / ncol)] * ncol)
    spacer(doc, 4)


def field_grid(doc, pairs, label_w=1.55, value_w=1.85):
    """pairs: list of ((label,value),(label,value)) — a 4-col label/value form."""
    t = doc.add_table(rows=len(pairs), cols=4); no_borders(t)
    for i, row in enumerate(pairs):
        cells = t.rows[i].cells
        for k, (label, value) in enumerate(row):
            lab = cells[k * 2]; val = cells[k * 2 + 1]
            shade(lab, C.TINT); cell_pad(lab, 90, top=14, bottom=14); cell_borders(lab, bottom=(4, C.WHITE))
            cell_text(lab, [(label, 8, C.NAVY, True)], line=1.0)
            cell_pad(val, 90, top=14, bottom=14); cell_borders(val, bottom=(6, C.TINT2))
            cell_text(val, [(value or "", 9, C.SLATE, False)], line=1.0)
    set_col_widths(t, [Inches(label_w), Inches(value_w), Inches(label_w), Inches(value_w)])
    spacer(doc, 2)


# ------------------------------------------------------------------ doc scaffold
def header_bottom_border(p, color, sz):
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), str(sz)); bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color.lstrip('#')); pbdr.append(bottom); pPr.append(pbdr)


def strip_trailing_empty(doc):
    body = doc.element.body
    for el in reversed(list(body)):
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p'):
            if not any(r.findall(qn('w:t')) for r in el.findall(qn('w:r'))):
                body.remove(el); continue
        break


def new_document(meta):
    """Create a branded doc with margins, logo header, and a doc-ref footer."""
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5); sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
    set_doc_default_font(doc, FONT)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].font.color.rgb = rgb(C.SLATE)
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(bs.logo_path('color'), width=Inches(1.35))
    header_bottom_border(hp, C.TINT, 8)
    doc_ref = meta.get("doc_ref", "BST-____")
    version = meta.get("version", "1.0")
    classification = meta.get("classification", "Confidential")
    fp = sec.footer.paragraphs[0]
    set_font(fp.add_run(f"{BRAND['website']}    ·    Doc Ref: {doc_ref}    ·    "
                        f"Version {version}    ·    Classification: {classification}"), 7.5, C.MUTED)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def finalize(doc, out_path):
    strip_trailing_empty(doc)
    anchor = doc.add_paragraph()
    apf = anchor.paragraph_format; apf.space_before = Pt(0); apf.space_after = Pt(0); apf.line_spacing = 1.0
    ar = anchor.add_run(" "); ar.font.size = Pt(1); ar.font.name = FONT
    _rfonts(ar._element.get_or_add_rPr(), FONT)
    d = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(d, exist_ok=True)
    doc.save(out_path)
    try:
        from embed_fonts import embed_fonts_in_docx
        embed_fonts_in_docx(out_path, [{
            "name": "Poppins",
            "regular": str(bs.FONTS / "Poppins-Regular.ttf"),
            "bold": str(bs.FONTS / "Poppins-Bold.ttf"),
        }])
    except Exception as e:
        sys.stderr.write(f"warning: font embed skipped ({e})\n")
    return out_path
