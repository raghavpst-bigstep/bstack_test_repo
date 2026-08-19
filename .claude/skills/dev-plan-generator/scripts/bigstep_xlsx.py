"""
BigStep-branded .xlsx toolkit (openpyxl) — reusable workbook composition helpers.

Renders a BigStep-branded multi-sheet workbook from a JSON content spec. Self-contained:
only needs `openpyxl` (no logo assets / Pillow). The brand palette matches the docx/PDF
engines (BigStep blue #1C62EC, Poppins font name; Excel falls back if Poppins is absent).

Deterministic + reproducible (standards P-2.2 / P-4.5): the model produces the CONTENT
(a spec), the script produces the WORKBOOK — no hand-built cells, colours, or formulas.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# palette (hex without '#') — matches bigstep_brand.C
class C:
    PRIMARY = "1C62EC"   # header fill
    NAVY    = "041342"   # title text on light
    SLATE   = "2B3D4F"   # body text
    TINT    = "E4F2FF"   # light panel
    TINT2   = "D4E6FF"   # zebra alt row
    WHITE   = "FFFFFF"
    MUTED   = "6B7A8D"
    RULE    = "B7CCEA"   # thin borders


FONT = "Poppins"
_thin = Side(style="thin", color=C.RULE)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def title_band(ws, row, text, ncols, size=14):
    """Merged blue title band across ncols columns."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(1, ncols))
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=size, bold=True, color=C.WHITE)
    c.fill = _fill(C.PRIMARY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 24
    return row + 1


def table(ws, row, headers, rows, widths=None, total=None):
    """Header row (blue) + zebra data rows with borders. Cell values starting with
    '=' are written as formulas.

    `total` (optional) renders a bold totals row with correct SUM ranges computed from
    the ACTUAL rendered positions — never trust spec-authored absolute cell refs, which
    break when a title band offsets the table:
        {"label": "Total", "label_col": 1, "sum_cols": [3]}   # 1-based column numbers
    Returns the next free row."""
    ncol = len(headers)
    # header
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color=C.WHITE)
        c.fill = _fill(C.PRIMARY)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        c.border = BORDER
    ws.row_dimensions[row].height = 20
    # body
    for i, r in enumerate(rows):
        rr = row + 1 + i
        fill = C.WHITE if i % 2 == 0 else C.TINT2
        for j in range(1, ncol + 1):
            val = r[j - 1] if j - 1 < len(r) else ""
            c = ws.cell(row=rr, column=j, value=val)   # openpyxl treats "=..." as a formula
            c.font = Font(name=FONT, size=9.5, color=C.SLATE)
            c.fill = _fill(fill)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
            c.border = BORDER
    next_row = row + 1 + len(rows)
    # totals row — SUM ranges computed from the real data-row span (row+1 .. row+len(rows))
    if total and rows:
        first, last = row + 1, row + len(rows)
        label_col = total.get("label_col", 1)
        lc = ws.cell(row=next_row, column=label_col, value=total.get("label", "Total"))
        lc.font = Font(name=FONT, size=9.5, bold=True, color=C.NAVY)
        lc.fill = _fill(C.TINT2); lc.border = BORDER
        lc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in total.get("sum_cols", []):
            L = get_column_letter(col)
            c = ws.cell(row=next_row, column=col, value=f"=SUM({L}{first}:{L}{last})")
            c.font = Font(name=FONT, size=9.5, bold=True, color=C.NAVY)
            c.fill = _fill(C.TINT2); c.border = BORDER
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        # fill remaining cells of the total row for a clean band
        for j in range(1, ncol + 1):
            if j == label_col or j in total.get("sum_cols", []):
                continue
            c = ws.cell(row=next_row, column=j); c.fill = _fill(C.TINT2); c.border = BORDER
        next_row += 1
    # widths
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    else:
        for j in range(1, ncol + 1):
            ws.column_dimensions[get_column_letter(j)].width = max(14, 90 // ncol)
    return next_row + 1


def note(ws, row, title, text, ncols=4, width_hint=None):
    """A wrapped note block: bold title row + a merged wrapped body cell."""
    if title:
        c = ws.cell(row=row, column=1, value=title)
        c.font = Font(name=FONT, size=11, bold=True, color=C.NAVY)
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(1, ncols))
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=9.5, color=C.SLATE)
    c.fill = _fill(C.TINT)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = max(28, 14 * (1 + len(str(text)) // (width_hint or 90)))
    c.border = BORDER
    return row + 2


def kv(ws, row, pairs):
    """Two-column label/value rows."""
    for label, value in pairs:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = Font(name=FONT, size=9.5, bold=True, color=C.NAVY)
        lc.fill = _fill(C.TINT)
        lc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        lc.border = BORDER
        vc = ws.cell(row=row, column=2, value=value)
        vc.font = Font(name=FONT, size=9.5, color=C.SLATE)
        vc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        vc.border = BORDER
        row += 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 26)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 40)
    return row + 1


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)   # start with no default sheet
    return wb


def add_sheet(wb, name):
    # Excel sheet names: max 31 chars, no []:*?/\
    safe = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Sheet"
    ws = wb.create_sheet(title=safe)
    ws.sheet_view.showGridLines = False
    return ws


def save(wb, out_path):
    import os
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path
