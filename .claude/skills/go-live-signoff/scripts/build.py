#!/usr/bin/env python3
"""
Build a BigStep-branded **Go-Live / Release Sign-Off** document (.docx) from an
optional JSON spec.

Usage:
    python build.py                       # blank, fillable form -> ./Go-Live-Release-Sign-Off.docx
    python build.py spec.json             # populated from spec  -> ./Go-Live-Release-Sign-Off.docx
    python build.py spec.json out.docx    # populated -> out.docx
    OUT=/path/out.docx python build.py spec.json

Every field is optional. Anything omitted renders as a blank line to write on,
so an empty spec reproduces the standard blank template exactly. The section
order is fixed and matches the source document:

    Title + purpose  ->  Release & Deployment Details  ->  Pre-Go-Live
    Readiness Checklist  ->  Rollback & Contingency  ->  Go / No-Go Decision
    ->  Sign-Off & Authorisation  ->  Doc-ref footer

The JSON schema is documented in ../references/content_spec.md.
"""
import sys, os, json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import bigstep_brand as bs
from bigstep_brand import C, BRAND
from embed_fonts import embed_fonts_in_docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Poppins"
BOX_EMPTY = "\u2610"   # ☐
BOX_CHECK = "\u2611"   # ☑


def rgb(h):
    return RGBColor(*bs.hx(h))


# ----------------------------------------------------------------- font plumbing
def _rfonts(rpr, name):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)


def set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    _rfonts(run._element.get_or_add_rPr(), FONT)


def set_doc_default_font(doc, name):
    el = doc.styles.element
    dd = el.find(qn('w:docDefaults'))
    if dd is None:
        dd = OxmlElement('w:docDefaults'); el.insert(0, dd)
    rprd = dd.find(qn('w:rPrDefault'))
    if rprd is None:
        rprd = OxmlElement('w:rPrDefault'); dd.append(rprd)
    rpr = rprd.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr'); rprd.append(rpr)
    _rfonts(rpr, name)
    n = doc.styles['Normal']; n.font.name = name
    _rfonts(n.element.get_or_add_rPr(), name)


# ----------------------------------------------------------------- cell helpers
def shade(cell, h):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), h.lstrip('#')); tcPr.append(sh)


def _set_border(edge_el, val, sz, color):
    edge_el.set(qn('w:val'), val)
    if val != 'nil':
        edge_el.set(qn('w:sz'), str(sz))
        edge_el.set(qn('w:space'), '0')
        edge_el.set(qn('w:color'), color.lstrip('#'))


def cell_borders(cell, top=None, left=None, bottom=None, right=None):
    """Each arg is None (nil) or a (sz, color) tuple."""
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None:
        tcPr.remove(existing)
    b = OxmlElement('w:tcBorders')
    for name, spec in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{name}')
        if spec is None:
            _set_border(e, 'nil', 0, C.WHITE)
        else:
            _set_border(e, 'single', spec[0], spec[1])
        b.append(e)
    tcPr.append(b)


def left_accent(cell, sz, color):
    cell_borders(cell, left=(sz, color))


def cell_pad(cell, dxa=140, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    vals = {'top': top if top is not None else dxa,
            'bottom': bottom if bottom is not None else dxa,
            'start': dxa, 'end': dxa}
    for m, v in vals.items():
        e = OxmlElement(f'w:{m}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcPr.append(mar)


def no_borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)


def set_col_widths(table, widths):
    """Force fixed column widths (python-docx needs it on every cell)."""
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = w


# ----------------------------------------------------------------- text helpers
def para(doc, runs, after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    for (t, s, c, b) in runs:
        set_font(p.add_run(t), s, c, b)
    return p


def cell_text(cell, runs, after=0, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.05, first=True):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    for (t, s, c, b) in runs:
        set_font(p.add_run(t), s, c, b)
    return p


def two_tone_title(doc, lead, tail, size=24, before=4, after=2):
    return para(doc, [(lead, size, C.NAVY, True), (tail, size, C.PRIMARY, True)],
                after=after, before=before)


def accent_rule(doc, width_in=1.4, color=None, sep=True):
    color = color or C.CYAN
    t = doc.add_table(rows=1, cols=1)
    set_col_widths(t, [Inches(width_in)])
    t.rows[0].height = Pt(2.2)
    no_borders(t); shade(t.rows[0].cells[0], color)
    r = t.rows[0].cells[0].paragraphs[0]; r.paragraph_format.space_after = Pt(0)
    set_font(r.add_run(" "), 2, color)
    if sep:
        # minimal separator so the next table doesn't merge with the rule table
        s = doc.add_paragraph()
        s.paragraph_format.space_after = Pt(1); s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.line_spacing = 1.0
        rr = s.add_run(" "); rr.font.size = Pt(2); rr.font.name = FONT
        _rfonts(rr._element.get_or_add_rPr(), FONT)


def spacer(doc, pts=6):
    doc.add_paragraph().paragraph_format.space_after = Pt(pts)


def section_head(doc, lead, tail, before=3):
    two_tone_title(doc, lead, tail, size=12, before=before, after=1)
    accent_rule(doc, 1.0)


# ----------------------------------------------------------------- form primitives
def field_grid(doc, pairs, label_w=1.55, value_w=1.75):
    """pairs: list of ((label_l, value_l), (label_r, value_r)) rows.
    Renders a 4-col form table: tinted label + underlined fill, x2 per row."""
    rows = len(pairs)
    t = doc.add_table(rows=rows, cols=4)
    widths = [Inches(label_w), Inches(value_w), Inches(label_w), Inches(value_w)]
    no_borders(t)
    for i, row in enumerate(pairs):
        cells = t.rows[i].cells
        for k, (label, value) in enumerate(row):
            lab = cells[k * 2]; val = cells[k * 2 + 1]
            shade(lab, C.TINT)
            cell_pad(lab, 90, top=14, bottom=14)
            cell_borders(lab, bottom=(4, C.WHITE))
            cell_text(lab, [(label, 8, C.NAVY, True)], line=1.0)
            cell_pad(val, 90, top=14, bottom=14)
            cell_borders(val, bottom=(6, C.TINT2))
            txt = value if value else ""
            cell_text(val, [(txt, 9, C.SLATE, False)], line=1.0)
    set_col_widths(t, widths)
    spacer(doc, 2)


def checklist_grid(doc, items, cols=2):
    """items: list of {text, checked}. 2-column zebra checkbox grid."""
    n = len(items)
    rows = (n + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    colw = Inches(6.7 / cols)
    no_borders(t)
    for idx, item in enumerate(items):
        r, cc = idx // cols, idx % cols
        cell = t.rows[r].cells[cc]
        fill = C.WHITE if (r % 2 == 0) else C.TINT
        shade(cell, fill)
        cell_pad(cell, 110, top=3, bottom=3)
        checked = bool(item.get("checked"))
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.88
        set_font(p.add_run((BOX_CHECK if checked else BOX_EMPTY) + "  "), 10, C.PRIMARY, True)
        set_font(p.add_run(item.get("text", "")), 7.5, C.SLATE, False)
    # fill any trailing empty cells in the last row with matching zebra shade
    used = n
    total = rows * cols
    for idx in range(used, total):
        r, cc = idx // cols, idx % cols
        cell = t.rows[r].cells[cc]
        shade(cell, C.WHITE if (r % 2 == 0) else C.TINT)
        cell_pad(cell, 110, top=3, bottom=3)
    set_col_widths(t, [colw] * cols)
    spacer(doc, 2)


def callout(doc, title, body_lines, accent=None, fill=None, width_in=6.7):
    fill = fill or C.TINT
    accent = accent or C.PRIMARY
    t = doc.add_table(rows=1, cols=1)
    set_col_widths(t, [Inches(width_in)])
    no_borders(t)
    cell = t.rows[0].cells[0]; shade(cell, fill); left_accent(cell, 34, accent); cell_pad(cell, 150, top=16, bottom=16)
    if title:
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.05
        set_font(p.add_run(title), 10, C.NAVY, True)
        for ln in body_lines:
            pp = cell.add_paragraph(); pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.1
            set_font(pp.add_run(ln), 9, C.SLATE, False)
    else:
        for i, ln in enumerate(body_lines):
            pp = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.1
            set_font(pp.add_run(ln), 9, C.SLATE, False)
    spacer(doc, 3)


def decision_row(doc, options, selected=None):
    """options: list of (key, label). selected: matching key or None."""
    t = doc.add_table(rows=1, cols=len(options))
    no_borders(t)
    colw = Inches(6.7 / len(options))
    for j, (key, label) in enumerate(options):
        cell = t.rows[0].cells[j]
        is_sel = (selected == key)
        shade(cell, C.TINT2 if is_sel else C.TINT)
        left_accent(cell, 28, C.PRIMARY if is_sel else C.CYAN)
        cell_pad(cell, 120, top=32, bottom=32)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.88
        set_font(p.add_run((BOX_CHECK if is_sel else BOX_EMPTY) + "  "), 12, C.PRIMARY, True)
        set_font(p.add_run(label), 9.5, C.NAVY, is_sel)
    set_col_widths(t, [colw] * len(options))
    spacer(doc, 3)


def fill_line(doc, label, value="", width_in=6.7, label_w=1.15):
    """A 'label: __________' write-on line as a 2-col borderless table."""
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    lab = t.rows[0].cells[0]; val = t.rows[0].cells[1]
    cell_pad(lab, 0, top=40, bottom=40); cell_pad(val, 0, top=40, bottom=40)
    cell_text(lab, [(label, 9, C.NAVY, True)], line=1.0)
    cell_borders(val, bottom=(6, C.TINT2))
    cell_text(val, [(value, 9, C.SLATE, False)], line=1.0)
    set_col_widths(t, [Inches(label_w), Inches(width_in - label_w)])
    spacer(doc, 3)


def signoff_block(doc, client, bigstep):
    """Two-column authorisation table with Name/Role/Signature/Date lines."""
    parties = [("Client / Customer Authorisation", client),
               ("BigStep Technologies Authorisation", bigstep)]
    t = doc.add_table(rows=1, cols=2)
    colw = Inches(3.35)
    no_borders(t)
    for j, (heading, data) in enumerate(parties):
        cell = t.rows[0].cells[j]
        cell_borders(cell,
                     top=(6, C.TINT2), left=(6, C.TINT2),
                     bottom=(6, C.TINT2), right=(6, C.TINT2))
        cell_pad(cell, 0)
        # header band inside the cell
        hp = cell.paragraphs[0]; hp.paragraph_format.space_after = Pt(4); hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.line_spacing = 1.0
        set_font(hp.add_run(heading), 9.5, C.NAVY, True)
        data = data or {}
        for pair in (("name", "role"), ("signature", "date")):
            lp = cell.add_paragraph(); lp.paragraph_format.space_after = Pt(2); lp.paragraph_format.line_spacing = 1.0
            for k, field in enumerate(pair):
                set_font(lp.add_run(("     " if k else "") + field.capitalize() + ":  "), 8.5, C.SLATE, True)
                v = data.get(field, "")
                set_font(lp.add_run(v if v else "\u00a0" * 9), 8.5, C.SLATE, False)
                _underline_run(lp.runs[-1])
    set_col_widths(t, [colw, colw])
    _sig_spacer_marker = True  # keep
    # shade header strip: give each cell's first paragraph a tinted background via table look is complex;
    # keep it clean with just the bordered box + bold heading.
    spacer(doc, 4)


def _underline_run(run):
    rpr = run._element.get_or_add_rPr()
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); u.set(qn('w:color'), C.TINT2.lstrip('#'))
    rpr.append(u)


def strip_trailing_empty(doc):
    """Remove trailing empty paragraphs before the final sectPr so a dangling
    paragraph after the last table doesn't spill onto a blank second page."""
    body = doc.element.body
    children = list(body)
    for el in reversed(children):
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p'):
            # empty if it has no text runs
            texts = el.findall(qn('w:r'))
            has_text = any(r.findall(qn('w:t')) for r in texts)
            if not has_text:
                body.remove(el)
                continue
        break


def header_bottom_border(p, color, sz):
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), str(sz)); bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color.lstrip('#'))
    pbdr.append(bottom); pPr.append(pbdr)


# ----------------------------------------------------------------- defaults
DEFAULT_DETAILS = [
    ("Project / Product", "project_product"), ("Client / Org", "client_org"),
    ("Release / Version", "release_version"), ("Release Type", "release_type"),
    ("BigStep Release Mgr", "bigstep_release_mgr"), ("Client Release Owner", "client_release_owner"),
    ("Target Environment", "target_environment"), ("Go-Live Date & Time", "go_live_datetime"),
    ("Deployment Window", "deployment_window"), ("Expected Downtime", "expected_downtime"),
]

DEFAULT_READINESS = [
    "UAT completed and signed off",
    "Critical / High defects resolved or accepted",
    "Code freeze; final build tagged & verified",
    "Release notes prepared & shared",
    "Deployment runbook documented & reviewed",
    "Database migration scripts tested",
    "Backups & rollback plan verified",
    "Production environment provisioned & access verified",
    "Security review completed (if applicable)",
    "Performance within agreed thresholds",
    "Monitoring, logging & alerting configured",
    "Integrations / dependencies confirmed ready",
    "Training & documentation delivered (if in scope)",
    "Support, on-call & escalation contacts confirmed",
]

DEFAULT_ROLLBACK = ("Trigger / criteria, method, estimated time, backup recovery point and decision "
                    "owner agreed and documented prior to deployment.")

DECISION_OPTIONS = [("go", "Go"),
                    ("conditional", "Conditional Go"),
                    ("nogo", "No-Go \u2014 postpone & remediate")]

DEFAULT_SUBTITLE = ("Confirms the release has met all readiness criteria and authorises "
                    "deployment to production.")


# ----------------------------------------------------------------- build
def build(spec, out_path):
    meta = spec.get("meta", {})
    details = spec.get("details", {})
    readiness = spec.get("readiness")
    rollback = spec.get("rollback", DEFAULT_ROLLBACK)
    decision = spec.get("decision", {})
    signoff = spec.get("signoff", {})

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5); sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
    set_doc_default_font(doc, FONT)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].font.color.rgb = rgb(C.SLATE)

    # header: color logo + tint rule
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(bs.logo_path('color'), width=Inches(1.35))
    header_bottom_border(hp, C.TINT, 8)
    # footer: website · doc ref · classification
    doc_ref = meta.get("doc_ref", "BST-GL-____")
    version = meta.get("version", "1.0")
    classification = meta.get("classification", "Confidential")
    fp = sec.footer.paragraphs[0]
    set_font(fp.add_run(f"{BRAND['website']}    \u00b7    Doc Ref: {doc_ref}    \u00b7    "
                        f"Version {version}    \u00b7    Classification: {classification}"), 7.5, C.MUTED)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # title + purpose
    two_tone_title(doc, "Go-Live / Release ", "Sign-Off", size=18, before=0, after=1)
    accent_rule(doc, 1.5)
    spacer(doc, 2)
    para(doc, [(meta.get("subtitle", DEFAULT_SUBTITLE), 9.5, C.SLATE, False)], after=4, line=1.15)

    # 1. Release & Deployment Details
    section_head(doc, "Release & Deployment ", "Details", before=3)
    pairs = []
    for i in range(0, len(DEFAULT_DETAILS), 2):
        (l_lab, l_key) = DEFAULT_DETAILS[i]
        (r_lab, r_key) = DEFAULT_DETAILS[i + 1]
        l_val = details.get(l_key, "")
        r_val = details.get(r_key, "")
        if r_key == "release_type" and not r_val:
            r_val = "Major / Minor / Patch / Hotfix"
        pairs.append(((l_lab, l_val), (r_lab, r_val)))
    field_grid(doc, pairs)

    # 2. Pre-Go-Live Readiness Checklist
    section_head(doc, "Pre-Go-Live Readiness ", "Checklist")
    if readiness:
        items = [{"text": r["text"], "checked": r.get("checked", False)} if isinstance(r, dict)
                 else {"text": str(r), "checked": False} for r in readiness]
    else:
        items = [{"text": t, "checked": False} for t in DEFAULT_READINESS]
    checklist_grid(doc, items)

    # 3. Rollback & Contingency
    section_head(doc, "Rollback & ", "Contingency")
    callout(doc, "Rollback plan", [rollback], accent=C.PRIMARY, fill=C.TINT)

    # 4. Go / No-Go Decision
    section_head(doc, "Go / No-Go Decision ", "(Select One)")
    decision_row(doc, DECISION_OPTIONS, selected=decision.get("selected"))
    fill_line(doc, "Conditions / notes:", decision.get("notes", ""))

    # 5. Sign-Off & Authorisation
    section_head(doc, "Sign-Off & ", "Authorisation")
    signoff_block(doc, signoff.get("client", {}), signoff.get("bigstep", {}))

    # A .docx body may not end on a table; renderers otherwise append a
    # full-height paragraph that spills to a blank page. Strip stray empties
    # and end on a deliberate 1pt anchor so the form stays a single page.
    strip_trailing_empty(doc)
    anchor = doc.add_paragraph()
    apf = anchor.paragraph_format
    apf.space_before = Pt(0); apf.space_after = Pt(0); apf.line_spacing = 1.0
    ar = anchor.add_run(" "); ar.font.size = Pt(1); ar.font.name = FONT
    _rfonts(ar._element.get_or_add_rPr(), FONT)

    doc.save(out_path)
    # embed Poppins so the brand holds without a local install
    embed_fonts_in_docx(out_path, [{
        "name": "Poppins",
        "regular": str(bs.FONTS / "Poppins-Regular.ttf"),
        "bold": str(bs.FONTS / "Poppins-Bold.ttf"),
    }])
    return out_path


def main():
    args = [a for a in sys.argv[1:]]
    spec = {}
    out = os.environ.get("OUT", "Go-Live-Release-Sign-Off.docx")
    if args:
        if args[0].endswith(".json"):
            with open(args[0]) as f:
                spec = json.load(f)
            if len(args) > 1:
                out = args[1]
        else:
            out = args[0]
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    path = build(spec, out)
    print("saved + fonts embedded:", path)


if __name__ == "__main__":
    main()
