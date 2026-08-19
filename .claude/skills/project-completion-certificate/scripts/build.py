#!/usr/bin/env python3
"""
Build a BigStep-branded **Project Completion Certificate** (.docx) from an
optional JSON spec.

Usage:
    python build.py                        # blank, fillable certificate
    python build.py spec.json              # populated from spec
    python build.py spec.json out.docx     # populated -> out.docx

Every field is optional. Anything omitted renders as a blank line to write on,
so an empty spec reproduces the standard blank certificate. Section order is
fixed and matches the source document:

    Certificate header + statement  ->  Project & Stakeholder Details  ->
    Deliverables & Milestones Completed  ->  Confirmation of Completion  ->
    Warranty, Support & Handover  ->  Acceptance & Authorisation

The JSON schema is documented in ../references/content_spec.md.
"""
import sys, os, json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import bigstep_brand as bs
from bigstep_brand import C, BRAND
from embed_fonts import embed_fonts_in_docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Poppins"
BOX_EMPTY = "\u2610"   # ☐
BOX_CHECK = "\u2611"   # ☑


def rgb(h):
    return RGBColor(*bs.hx(h))


# ----------------------------------------------------------------- font plumbing
def _rfonts(rpr, name):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)


def set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    _rfonts(run._element.get_or_add_rPr(), FONT)


def set_doc_default_font(doc, name):
    el = doc.styles.element
    dd = el.find(qn('w:docDefaults'))
    if dd is None:
        dd = OxmlElement('w:docDefaults'); el.insert(0, dd)
    rprd = dd.find(qn('w:rPrDefault'))
    if rprd is None:
        rprd = OxmlElement('w:rPrDefault'); dd.append(rprd)
    rpr = rprd.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr'); rprd.append(rpr)
    _rfonts(rpr, name)
    n = doc.styles['Normal']; n.font.name = name
    _rfonts(n.element.get_or_add_rPr(), name)


# ----------------------------------------------------------------- cell helpers
def shade(cell, h):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), h.lstrip('#')); tcPr.append(sh)


def _set_border(edge_el, val, sz, color):
    edge_el.set(qn('w:val'), val)
    if val != 'nil':
        edge_el.set(qn('w:sz'), str(sz))
        edge_el.set(qn('w:space'), '0')
        edge_el.set(qn('w:color'), color.lstrip('#'))


def cell_borders(cell, top=None, left=None, bottom=None, right=None):
    """Each arg is None (nil) or a (sz, color) tuple."""
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None:
        tcPr.remove(existing)
    b = OxmlElement('w:tcBorders')
    for name, spec in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{name}')
        if spec is None:
            _set_border(e, 'nil', 0, C.WHITE)
        else:
            _set_border(e, 'single', spec[0], spec[1])
        b.append(e)
    tcPr.append(b)


def left_accent(cell, sz, color):
    cell_borders(cell, left=(sz, color))


def cell_pad(cell, dxa=140, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    vals = {'top': top if top is not None else dxa,
            'bottom': bottom if bottom is not None else dxa,
            'start': dxa, 'end': dxa}
    for m, v in vals.items():
        e = OxmlElement(f'w:{m}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcPr.append(mar)


def no_borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)


def set_col_widths(table, widths):
    """Force fixed column widths (python-docx needs it on every cell)."""
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = w


# ----------------------------------------------------------------- text helpers
def para(doc, runs, after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    for (t, s, c, b) in runs:
        set_font(p.add_run(t), s, c, b)
    return p


def cell_text(cell, runs, after=0, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.05, first=True):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    for (t, s, c, b) in runs:
        set_font(p.add_run(t), s, c, b)
    return p


def two_tone_title(doc, lead, tail, size=24, before=4, after=2):
    return para(doc, [(lead, size, C.NAVY, True), (tail, size, C.PRIMARY, True)],
                after=after, before=before)


def accent_rule(doc, width_in=1.4, color=None, sep=True):
    color = color or C.CYAN
    t = doc.add_table(rows=1, cols=1)
    set_col_widths(t, [Inches(width_in)])
    t.rows[0].height = Pt(2.2)
    no_borders(t); shade(t.rows[0].cells[0], color)
    r = t.rows[0].cells[0].paragraphs[0]; r.paragraph_format.space_after = Pt(0)
    set_font(r.add_run(" "), 2, color)
    if sep:
        # minimal separator so the next table doesn't merge with the rule table
        s = doc.add_paragraph()
        s.paragraph_format.space_after = Pt(1); s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.line_spacing = 1.0
        rr = s.add_run(" "); rr.font.size = Pt(2); rr.font.name = FONT
        _rfonts(rr._element.get_or_add_rPr(), FONT)


def spacer(doc, pts=6):
    doc.add_paragraph().paragraph_format.space_after = Pt(pts)


def section_head(doc, lead, tail, before=3):
    two_tone_title(doc, lead, tail, size=12, before=before, after=1)
    accent_rule(doc, 1.0)


# ----------------------------------------------------------------- form primitives
def field_grid(doc, pairs, label_w=1.55, value_w=1.75):
    """pairs: list of ((label_l, value_l), (label_r, value_r)) rows.
    Renders a 4-col form table: tinted label + underlined fill, x2 per row."""
    rows = len(pairs)
    t = doc.add_table(rows=rows, cols=4)
    widths = [Inches(label_w), Inches(value_w), Inches(label_w), Inches(value_w)]
    no_borders(t)
    for i, row in enumerate(pairs):
        cells = t.rows[i].cells
        for k, (label, value) in enumerate(row):
            lab = cells[k * 2]; val = cells[k * 2 + 1]
            shade(lab, C.TINT)
            cell_pad(lab, 90, top=14, bottom=14)
            cell_borders(lab, bottom=(4, C.WHITE))
            cell_text(lab, [(label, 8, C.NAVY, True)], line=1.0)
            cell_pad(val, 90, top=14, bottom=14)
            cell_borders(val, bottom=(6, C.TINT2))
            txt = value if value else ""
            cell_text(val, [(txt, 9, C.SLATE, False)], line=1.0)
    set_col_widths(t, widths)
    spacer(doc, 2)


def checklist_grid(doc, items, cols=2):
    """items: list of {text, checked}. 2-column zebra checkbox grid."""
    n = len(items)
    rows = (n + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    colw = Inches(6.7 / cols)
    no_borders(t)
    for idx, item in enumerate(items):
        r, cc = idx // cols, idx % cols
        cell = t.rows[r].cells[cc]
        fill = C.WHITE if (r % 2 == 0) else C.TINT
        shade(cell, fill)
        cell_pad(cell, 110, top=3, bottom=3)
        checked = bool(item.get("checked"))
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.88
        set_font(p.add_run((BOX_CHECK if checked else BOX_EMPTY) + "  "), 10, C.PRIMARY, True)
        set_font(p.add_run(item.get("text", "")), 7.5, C.SLATE, False)
    # fill any trailing empty cells in the last row with matching zebra shade
    used = n
    total = rows * cols
    for idx in range(used, total):
        r, cc = idx // cols, idx % cols
        cell = t.rows[r].cells[cc]
        shade(cell, C.WHITE if (r % 2 == 0) else C.TINT)
        cell_pad(cell, 110, top=3, bottom=3)
    set_col_widths(t, [colw] * cols)
    spacer(doc, 2)


def callout(doc, title, body_lines, accent=None, fill=None, width_in=6.7):
    fill = fill or C.TINT
    accent = accent or C.PRIMARY
    t = doc.add_table(rows=1, cols=1)
    set_col_widths(t, [Inches(width_in)])
    no_borders(t)
    cell = t.rows[0].cells[0]; shade(cell, fill); left_accent(cell, 34, accent); cell_pad(cell, 150, top=16, bottom=16)
    if title:
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.05
        set_font(p.add_run(title), 10, C.NAVY, True)
        for ln in body_lines:
            pp = cell.add_paragraph(); pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.1
            set_font(pp.add_run(ln), 9, C.SLATE, False)
    else:
        for i, ln in enumerate(body_lines):
            pp = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.1
            set_font(pp.add_run(ln), 9, C.SLATE, False)
    spacer(doc, 3)


def decision_row(doc, options, selected=None):
    """options: list of (key, label). selected: matching key or None."""
    t = doc.add_table(rows=1, cols=len(options))
    no_borders(t)
    colw = Inches(6.7 / len(options))
    for j, (key, label) in enumerate(options):
        cell = t.rows[0].cells[j]
        is_sel = (selected == key)
        shade(cell, C.TINT2 if is_sel else C.TINT)
        left_accent(cell, 28, C.PRIMARY if is_sel else C.CYAN)
        cell_pad(cell, 120, top=32, bottom=32)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.88
        set_font(p.add_run((BOX_CHECK if is_sel else BOX_EMPTY) + "  "), 12, C.PRIMARY, True)
        set_font(p.add_run(label), 9.5, C.NAVY, is_sel)
    set_col_widths(t, [colw] * len(options))
    spacer(doc, 3)


def fill_line(doc, label, value="", width_in=6.7, label_w=1.15):
    """A 'label: __________' write-on line as a 2-col borderless table."""
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    lab = t.rows[0].cells[0]; val = t.rows[0].cells[1]
    cell_pad(lab, 0, top=40, bottom=40); cell_pad(val, 0, top=40, bottom=40)
    cell_text(lab, [(label, 9, C.NAVY, True)], line=1.0)
    cell_borders(val, bottom=(6, C.TINT2))
    cell_text(val, [(value, 9, C.SLATE, False)], line=1.0)
    set_col_widths(t, [Inches(label_w), Inches(width_in - label_w)])
    spacer(doc, 3)


def signoff_block(doc, client, bigstep, headings=None):
    """Two-column authorisation table with Name/Role/Signature/Date lines."""
    headings = headings or ("Client / Customer Authorisation",
                            "BigStep Technologies Authorisation")
    parties = [(headings[0], client), (headings[1], bigstep)]
    t = doc.add_table(rows=1, cols=2)
    colw = Inches(3.35)
    no_borders(t)
    for j, (heading, data) in enumerate(parties):
        cell = t.rows[0].cells[j]
        cell_borders(cell,
                     top=(6, C.TINT2), left=(6, C.TINT2),
                     bottom=(6, C.TINT2), right=(6, C.TINT2))
        cell_pad(cell, 0)
        # header band inside the cell
        hp = cell.paragraphs[0]; hp.paragraph_format.space_after = Pt(4); hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.line_spacing = 1.0
        set_font(hp.add_run(heading), 9.5, C.NAVY, True)
        data = data or {}
        for pair in (("name", "role"), ("signature", "date")):
            lp = cell.add_paragraph(); lp.paragraph_format.space_after = Pt(2); lp.paragraph_format.line_spacing = 1.0
            for k, field in enumerate(pair):
                set_font(lp.add_run(("     " if k else "") + field.capitalize() + ":  "), 8.5, C.SLATE, True)
                v = data.get(field, "")
                set_font(lp.add_run(v if v else "\u00a0" * 9), 8.5, C.SLATE, False)
                _underline_run(lp.runs[-1])
    set_col_widths(t, [colw, colw])
    _sig_spacer_marker = True  # keep
    # shade header strip: give each cell's first paragraph a tinted background via table look is complex;
    # keep it clean with just the bordered box + bold heading.
    spacer(doc, 4)


def _underline_run(run):
    rpr = run._element.get_or_add_rPr()
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); u.set(qn('w:color'), C.TINT2.lstrip('#'))
    rpr.append(u)


def strip_trailing_empty(doc):
    """Remove trailing empty paragraphs before the final sectPr so a dangling
    paragraph after the last table doesn't spill onto a blank second page."""
    body = doc.element.body
    children = list(body)
    for el in reversed(children):
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p'):
            # empty if it has no text runs
            texts = el.findall(qn('w:r'))
            has_text = any(r.findall(qn('w:t')) for r in texts)
            if not has_text:
                body.remove(el)
                continue
        break


def header_bottom_border(p, color, sz):
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), str(sz)); bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color.lstrip('#'))
    pbdr.append(bottom); pPr.append(pbdr)


def center_kicker(doc, text):
    p = para(doc, [(text, 10, C.CYAN, True)], after=1, before=0,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    # letter-spacing for the uppercase kicker
    for r in p.runs:
        rpr = r._element.get_or_add_rPr()
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '60'); rpr.append(sp)
    return p


def two_tone_center(doc, lead, tail, size=26, before=2, after=2):
    return para(doc, [(lead, size, C.NAVY, True), (tail, size, C.PRIMARY, True)],
                after=after, before=before, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)


def center_rule(doc, width_in=1.8, color=None):
    color = color or C.CYAN
    t = doc.add_table(rows=1, cols=1)
    set_col_widths(t, [Inches(width_in)])
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].height = Pt(2.4)
    no_borders(t); shade(t.rows[0].cells[0], color)
    r = t.rows[0].cells[0].paragraphs[0]
    r.alignment = WD_ALIGN_PARAGRAPH.CENTER; r.paragraph_format.space_after = Pt(0)
    set_font(r.add_run(" "), 2, color)
    s = doc.add_paragraph()
    s.paragraph_format.space_after = Pt(1); s.paragraph_format.space_before = Pt(0)
    rr = s.add_run(" "); rr.font.size = Pt(2)


def deliverables_table(doc, rows, count=3):
    """Branded table: # | Deliverable / Milestone | Status | Date Delivered.
    `rows` is a list of dicts {n, deliverable, status, date}; if empty, render
    `count` blank numbered rows to fill in."""
    headers = ["#", "Deliverable / Milestone", "Status", "Date Delivered"]
    widths = [Inches(0.4), Inches(3.7), Inches(1.3), Inches(1.3)]
    bt = doc.add_table(rows=1, cols=4)
    for j, h in enumerate(headers):
        c = bt.rows[0].cells[j]
        shade(c, C.PRIMARY); cell_pad(c, 110, top=40, bottom=40)
        align = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
        p = c.paragraphs[0]; p.alignment = align; p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(h), 9, C.WHITE, True)
    body = rows if rows else [{"n": i + 1} for i in range(count)]
    for i, row in enumerate(body):
        cells = bt.add_row().cells
        fill = C.WHITE if i % 2 == 0 else C.TINT2
        vals = [str(row.get("n", i + 1)), row.get("deliverable", ""),
                row.get("status", ""), row.get("date", "")]
        for j, val in enumerate(vals):
            c = cells[j]; shade(c, fill); cell_pad(c, 110, top=60, bottom=60)
            align = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            p = c.paragraphs[0]; p.alignment = align; p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(val), 9, C.SLATE if j else C.NAVY, j == 0)
    set_col_widths(bt, widths)
    # thin grid
    tblPr = bt._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4'); e.set(qn('w:color'), C.TINT.lstrip('#')); b.append(e)
    tblPr.append(b)
    spacer(doc, 3)


# ----------------------------------------------------------------- defaults
DEFAULT_DETAILS = [
    ("Project / Product", "project_product"), ("Client / Org", "client_org"),
    ("Project / Contract Ref.", "project_ref"), ("Certificate No.", "certificate_no"),
    ("BigStep PM", "bigstep_pm"), ("Client Sponsor", "client_sponsor"),
    ("Start Date", "start_date"), ("Completion Date", "completion_date"),
]

DEFAULT_WARRANTY = [
    ("Warranty / Defect-Liability", "warranty_liability"), ("Warranty Start & End", "warranty_dates"),
    ("Post-Completion Support", "post_completion_support"), ("Support / Escalation", "support_escalation"),
    ("Outstanding / Known Issues", "outstanding_issues"), ("Final Payment Status", "final_payment_status"),
]

DEFAULT_CONFIRMATION = ("All in-scope deliverables handed over; UAT signed off; solution deployed "
                        "& operational; all Critical / High defects resolved; documentation, source "
                        "code & knowledge transfer provided per scope.")

DEFAULT_STATEMENT = ("This certifies that the project described below has been successfully completed "
                     "and delivered by BigStep Technologies in accordance with the agreed scope, "
                     "requirements and acceptance criteria, and has been formally accepted by the Client.")


def _grid_pairs(defs, data, defaults=None):
    defaults = defaults or {}
    pairs = []
    for i in range(0, len(defs), 2):
        (l_lab, l_key) = defs[i]
        (r_lab, r_key) = defs[i + 1]
        l_val = data.get(l_key, "") or defaults.get(l_key, "")
        r_val = data.get(r_key, "") or defaults.get(r_key, "")
        pairs.append(((l_lab, l_val), (r_lab, r_val)))
    return pairs


# ----------------------------------------------------------------- build
def build(spec, out_path):
    meta = spec.get("meta", {})
    details = spec.get("details", {})
    deliverables = spec.get("deliverables") or []
    confirmation = spec.get("confirmation", DEFAULT_CONFIRMATION)
    warranty = spec.get("warranty", {})
    signoff = spec.get("signoff", {})

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5); sec.bottom_margin = Inches(0.42)
    sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
    set_doc_default_font(doc, FONT)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].font.color.rgb = rgb(C.SLATE)

    # header: color logo + tint rule
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(bs.logo_path('color'), width=Inches(1.35))
    header_bottom_border(hp, C.TINT, 8)
    # footer: version · classification · issued by
    version = meta.get("version", "1.0")
    classification = meta.get("classification", "Confidential")
    fp = sec.footer.paragraphs[0]
    set_font(fp.add_run(f"Version {version}    \u00b7    Classification: {classification}"
                        f"    \u00b7    Issued by {BRAND['company']}"), 7.5, C.MUTED)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # certificate header (centered)
    center_kicker(doc, "CERTIFICATE OF")
    two_tone_center(doc, "Project ", "Completion", size=26, before=0, after=2)
    center_rule(doc, 1.9)
    para(doc, [(meta.get("statement", DEFAULT_STATEMENT), 9.5, C.SLATE, False)],
         after=6, before=1, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2)

    # 1. Project & Stakeholder Details
    section_head(doc, "Project & Stakeholder ", "Details", before=3)
    field_grid(doc, _grid_pairs(DEFAULT_DETAILS, details,
                                defaults={"certificate_no": "BST-PCC-____"}))

    # 2. Deliverables & Milestones Completed
    section_head(doc, "Deliverables & Milestones ", "Completed")
    deliverables_table(doc, deliverables, count=3)

    # 3. Confirmation of Completion
    section_head(doc, "Confirmation of ", "Completion")
    callout(doc, "As of the completion date", [confirmation], accent=C.PRIMARY, fill=C.TINT)

    # 4. Warranty, Support & Handover
    section_head(doc, "Warranty, Support & ", "Handover")
    field_grid(doc, _grid_pairs(DEFAULT_WARRANTY, warranty))

    # 5. Acceptance & Authorisation
    section_head(doc, "Acceptance & ", "Authorisation")
    signoff_block(doc, signoff.get("client", {}), signoff.get("bigstep", {}),
                  headings=("Client / Customer Acceptance", "BigStep Technologies Authorisation"))

    # A .docx body may not end on a table; end on a 1pt anchor so a dangling
    # full-height paragraph doesn't spill onto a blank page.
    strip_trailing_empty(doc)
    anchor = doc.add_paragraph()
    apf = anchor.paragraph_format
    apf.space_before = Pt(0); apf.space_after = Pt(0); apf.line_spacing = 1.0
    ar = anchor.add_run(" "); ar.font.size = Pt(1); ar.font.name = FONT
    _rfonts(ar._element.get_or_add_rPr(), FONT)

    doc.save(out_path)
    embed_fonts_in_docx(out_path, [{
        "name": "Poppins",
        "regular": str(bs.FONTS / "Poppins-Regular.ttf"),
        "bold": str(bs.FONTS / "Poppins-Bold.ttf"),
    }])
    return out_path


def main():
    args = [a for a in sys.argv[1:]]
    spec = {}
    out = os.environ.get("OUT", "Project-Completion-Certificate.docx")
    if args:
        if args[0].endswith(".json"):
            with open(args[0]) as f:
                spec = json.load(f)
            if len(args) > 1:
                out = args[1]
        else:
            out = args[0]
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    path = build(spec, out)
    print("saved + fonts embedded:", path)


if __name__ == "__main__":
    main()
